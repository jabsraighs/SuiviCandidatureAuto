# -*- coding: utf-8 -*-
"""
excel_offres.py — Ajoute les nouvelles offres à un fichier Excel cumulatif
(offres.xlsx), sans jamais écraser les lignes déjà présentes.

Colonnes : Date découverte | Titre | Source | Lien | Statut
La colonne "Statut" est laissée vide pour que tu la remplisses toi-même
("à regarder", "postulé", "pas intéressé"...).
"""

import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ENTETES = ["Date découverte", "Titre", "Source", "Lien", "Statut"]
POLICE = "Arial"


def _creer_classeur(chemin):
    wb = Workbook()
    ws = wb.active
    ws.title = "Offres"
    for col_idx, entete in enumerate(ENTETES, start=1):
        cell = ws.cell(row=1, column=col_idx, value=entete)
        cell.font = Font(name=POLICE, bold=True)
    ws.freeze_panes = "A2"
    largeurs = [16, 55, 30, 60, 16]
    for col_idx, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = largeur
    wb.save(chemin)
    return wb


def ajouter_offres(chemin_xlsx, offres):
    """Ajoute les offres en haut du fichier (juste sous l'en-tête).
    Les plus récentes apparaissent donc toujours en premier.
    Crée le fichier s'il n'existe pas."""
    if not offres:
        return 0

    if os.path.exists(chemin_xlsx):
        wb = load_workbook(chemin_xlsx)
        ws = wb["Offres"] if "Offres" in wb.sheetnames else wb.active
    else:
        wb = _creer_classeur(chemin_xlsx)
        ws = wb["Offres"]

    horodatage = datetime.now().strftime("%Y-%m-%d %H:%M")

    # On insère autant de lignes vides que d'offres, juste après l'en-tête (ligne 1)
    ws.insert_rows(2, amount=len(offres))

    ligne = 2
    for offre in offres:
        valeurs = [
            horodatage,
            offre.get("titre", ""),
            offre.get("recherche", ""),
            offre.get("lien", ""),
            "",  # Statut, à remplir manuellement
        ]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cell = ws.cell(row=ligne, column=col_idx, value=valeur)
            cell.font = Font(name=POLICE)
        ligne += 1

    wb.save(chemin_xlsx)
    return len(offres)